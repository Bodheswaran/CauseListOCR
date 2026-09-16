"""
Output formatting for CauseListOCR
Export extracted data to various formats (CSV, JSON, Excel)
"""

import json
import csv
import logging
from typing import List, Dict, Optional
from io import StringIO
from datetime import datetime

logger = logging.getLogger(__name__)


class OutputFormatter:
    """Format and export extracted records."""
    
    @staticmethod
    def to_json(
        records: List[Dict],
        pretty: bool = True,
        include_metadata: bool = True
    ) -> str:
        """Export records as JSON.
        
        Args:
            records: List of extracted records
            pretty: Pretty-print JSON
            include_metadata: Include export metadata
            
        Returns:
            JSON string
        """
        output = {}
        
        if include_metadata:
            output["metadata"] = {
                "exported_at": datetime.now().isoformat(),
                "total_records": len(records),
                "format_version": "1.0"
            }
        
        output["records"] = records
        
        indent = 2 if pretty else None
        json_str = json.dumps(output, indent=indent, default=str)
        
        logger.info(f"Exported {len(records)} records to JSON")
        return json_str
    
    @staticmethod
    def to_csv(
        records: List[Dict],
        fieldnames: Optional[List[str]] = None,
        include_all_fields: bool = False
    ) -> str:
        """Export records as CSV.
        
        Args:
            records: List of extracted records
            fieldnames: List of field names to include (default: standard fields)
            include_all_fields: Include all fields found in records
            
        Returns:
            CSV string
        """
        if not records:
            return ""
        
        # Determine fields to export
        if fieldnames:
            csv_fieldnames = fieldnames
        elif include_all_fields:
            # Collect all unique field names
            all_fields = set()
            for record in records:
                all_fields.update(record.keys())
            csv_fieldnames = sorted(list(all_fields))
        else:
            # Standard fields
            csv_fieldnames = [
                "case_number",
                "hearing_date",
                "party_name",
                "advocate",
                "case_number_confidence",
                "date_confidence",
                "party_name_confidence",
                "advocate_confidence",
                "overall_confidence",
                "confidence_level",
                "needs_review",
                "is_duplicate",
                "validation_issues"
            ]
        
        output = StringIO()
        writer = csv.DictWriter(
            output,
            fieldnames=csv_fieldnames,
            restval="-",
            extrasaction="ignore"
        )
        
        writer.writeheader()
        for record in records:
            # Convert lists to strings for CSV
            csv_record = {}
            for field in csv_fieldnames:
                value = record.get(field, "-")
                if isinstance(value, list):
                    csv_record[field] = "; ".join(str(v) for v in value)
                else:
                    csv_record[field] = value
            writer.writerow(csv_record)
        
        csv_str = output.getvalue()
        logger.info(f"Exported {len(records)} records to CSV")
        return csv_str
    
    @staticmethod
    def to_excel(
        records: List[Dict],
        filename: str = "cause_list.xlsx",
        fieldnames: Optional[List[str]] = None
    ) -> bytes:
        """Export records as Excel file.
        
        Args:
            records: List of extracted records
            filename: Output filename
            fieldnames: List of field names to include
            
        Returns:
            Excel file bytes
        """
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            logger.error("openpyxl not installed. Install with: pip install openpyxl")
            raise ImportError("openpyxl required for Excel export")
        
        # Determine fields
        if fieldnames is None:
            fieldnames = [
                "case_number",
                "hearing_date",
                "party_name",
                "advocate",
                "overall_confidence",
                "confidence_level",
                "needs_review"
            ]
        
        # Create workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Cause List"
        
        # Write header
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        
        for col_idx, field in enumerate(fieldnames, 1):
            cell = ws.cell(row=1, column=col_idx, value=field)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Write data
        for row_idx, record in enumerate(records, 2):
            for col_idx, field in enumerate(fieldnames, 1):
                value = record.get(field, "-")
                if isinstance(value, list):
                    value = "; ".join(str(v) for v in value)
                
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        
        # Auto-adjust column widths
        for col_idx, field in enumerate(fieldnames, 1):
            max_length = len(field)
            for row_idx in range(2, len(records) + 2):
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            
            adjusted_width = min(max_length + 2, 50)  # Cap at 50
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = adjusted_width
        
        # Save to bytes
        from io import BytesIO
        output = BytesIO()
        wb.save(output)
        excel_bytes = output.getvalue()
        
        logger.info(f"Exported {len(records)} records to Excel ({filename})")
        return excel_bytes
    
    @staticmethod
    def to_html(
        records: List[Dict],
        title: str = "Cause List",
        include_confidence: bool = True
    ) -> str:
        """Export records as HTML table.
        
        Args:
            records: List of extracted records
            title: Page title
            include_confidence: Include confidence columns
            
        Returns:
            HTML string
        """
        html = f"""<!DOCTYPE html>
<html>
<head>
    <meta charset="UTF-8">
    <title>{title}</title>
    <style>
        body {{ font-family: Arial, sans-serif; margin: 20px; }}
        h1 {{ color: #333; }}
        table {{ border-collapse: collapse; width: 100%; }}
        th {{ background-color: #4472C4; color: white; padding: 10px; text-align: left; }}
        td {{ border: 1px solid #ddd; padding: 8px; }}
        tr:nth-child(even) {{ background-color: #f9f9f9; }}
        tr:hover {{ background-color: #f5f5f5; }}
        .high {{ color: green; }}
        .medium {{ color: orange; }}
        .low {{ color: red; }}
        .critical {{ color: darkred; font-weight: bold; }}
        .review {{ background-color: #fff3cd; }}
    </style>
</head>
<body>
<h1>{title}</h1>
<p>Exported: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Total Records: {len(records)}</p>
<table>
<thead>
<tr>
    <th>Case Number</th>
    <th>Hearing Date</th>
    <th>Party Name</th>
    <th>Advocate</th>
"""
        
        if include_confidence:
            html += """    <th>Confidence</th>
    <th>Level</th>
    <th>Needs Review</th>
"""
        
        html += """</tr>
</thead>
<tbody>
"""
        
        for record in records:
            needs_review = record.get('needs_review', False)
            row_class = " class='review'" if needs_review else ""
            conf_level = record.get('confidence_level', 'UNKNOWN').lower()
            
            html += f"<tr{row_class}>\n"
            html += f"    <td>{record.get('case_number', '-')}</td>\n"
            html += f"    <td>{record.get('hearing_date', '-')}</td>\n"
            html += f"    <td>{record.get('party_name', '-')}</td>\n"
            html += f"    <td>{record.get('advocate', '-')}</td>\n"
            
            if include_confidence:
                conf = record.get('overall_confidence', 0)
                html += f"    <td>{conf:.2%}</td>\n"
                html += f"    <td><span class='{conf_level}'>{conf_level.upper()}</span></td>\n"
                html += f"    <td>{'Yes' if needs_review else 'No'}</td>\n"
            
            html += "</tr>\n"
        
        html += """</tbody>
</table>
</body>
</html>
"""
        
        logger.info(f"Exported {len(records)} records to HTML")
        return html
    
    @staticmethod
    def get_export_summary(records: List[Dict]) -> Dict:
        """Get summary statistics for records.
        
        Args:
            records: List of extracted records
            
        Returns:
            Dictionary with summary statistics
        """
        total = len(records)
        needs_review = sum(1 for r in records if r.get('needs_review', False))
        high_conf = sum(1 for r in records if r.get('confidence_level') == 'HIGH')
        medium_conf = sum(1 for r in records if r.get('confidence_level') == 'MEDIUM')
        low_conf = sum(1 for r in records if r.get('confidence_level') == 'LOW')
        critical = sum(1 for r in records if r.get('confidence_level') == 'CRITICAL')
        duplicates = sum(1 for r in records if r.get('is_duplicate', False))
        
        avg_confidence = sum(r.get('overall_confidence', 0) for r in records) / total if total > 0 else 0
        
        return {
            "total_records": total,
            "records_needing_review": needs_review,
            "average_confidence": avg_confidence,
            "confidence_breakdown": {
                "HIGH": high_conf,
                "MEDIUM": medium_conf,
                "LOW": low_conf,
                "CRITICAL": critical
            },
            "duplicates_found": duplicates
        }
